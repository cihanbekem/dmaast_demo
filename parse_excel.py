import openpyxl
import json
import os
import sys

DATA_DIR = os.path.join(os.path.dirname(__file__), "data")
OUTPUT_DIR = os.path.join(os.path.dirname(__file__), "src", "data")

def read_sheet(wb, sheet_name):
    """Read a sheet into a list of dicts using the first row as headers."""
    if sheet_name not in wb.sheetnames:
        for sn in wb.sheetnames:
            if sheet_name.lower() in sn.lower():
                sheet_name = sn
                break
        else:
            return []
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h).strip() if h else f"col_{i}" for i, h in enumerate(rows[0])]
    data = []
    for row in rows[1:]:
        if all(v is None for v in row):
            continue
        record = {}
        for i, val in enumerate(row):
            if i < len(headers):
                record[headers[i]] = val
        data.append(record)
    return data

def parse_file(filepath):
    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    result = {
        "sheets": wb.sheetnames,
        "hierarchy": [],
        "entity_mapping": [],
        "process_mapping": [],
        "bom_relation_nodes": [],
        "relations": [],
    }

    for sn in wb.sheetnames:
        sn_lower = sn.lower()
        if "hierarchy" in sn_lower or "product_hierarchy" in sn_lower:
            result["hierarchy"] = read_sheet(wb, sn)
        elif "entity" in sn_lower:
            result["entity_mapping"] = read_sheet(wb, sn)
        elif "process" in sn_lower:
            result["process_mapping"] = read_sheet(wb, sn)
        elif "bom" in sn_lower and "relation" in sn_lower:
            result["bom_relation_nodes"] = read_sheet(wb, sn)
        elif "relation" in sn_lower and "bom" not in sn_lower:
            result["relations"] = read_sheet(wb, sn)

    wb.close()
    return result

def sanitize(obj):
    """Make JSON-serializable."""
    if isinstance(obj, dict):
        return {k: sanitize(v) for k, v in obj.items()}
    if isinstance(obj, list):
        return [sanitize(v) for v in obj]
    if isinstance(obj, (int, float, str, bool)) or obj is None:
        return obj
    return str(obj)

def main():
    os.makedirs(OUTPUT_DIR, exist_ok=True)

    all_data = {}
    files = sorted(os.listdir(DATA_DIR))

    for fname in files:
        if not fname.endswith(".xlsx"):
            continue
        fpath = os.path.join(DATA_DIR, fname)
        print(f"Parsing: {fname}")

        key = fname.replace("DES_ProductHierarchy_", "").replace(".xlsx", "")
        parsed = parse_file(fpath)
        all_data[key] = sanitize(parsed)

        print(f"  Sheets: {parsed['sheets']}")
        print(f"  Hierarchy rows: {len(parsed['hierarchy'])}")
        print(f"  Entity Mapping rows: {len(parsed['entity_mapping'])}")
        print(f"  Process Mapping rows: {len(parsed['process_mapping'])}")
        print(f"  BOM Relation Nodes rows: {len(parsed['bom_relation_nodes'])}")
        print(f"  Relations rows: {len(parsed['relations'])}")

    # Write combined JSON
    out_path = os.path.join(OUTPUT_DIR, "hierarchyData.json")
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(all_data, f, ensure_ascii=False, indent=2)
    print(f"\nWrote combined data to {out_path}")

    # Also write a summary
    summary = {}
    for key, data in all_data.items():
        h = data["hierarchy"]
        if h:
            cols = list(h[0].keys()) if h else []
            summary[key] = {
                "total_rows": len(h),
                "columns": cols,
                "sample_row": h[0] if h else None
            }
    summary_path = os.path.join(OUTPUT_DIR, "summary.json")
    with open(summary_path, "w", encoding="utf-8") as f:
        json.dump(summary, f, ensure_ascii=False, indent=2)
    print(f"Wrote summary to {summary_path}")

if __name__ == "__main__":
    main()
